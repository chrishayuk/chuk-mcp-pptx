"""
Chart utilities for reusable chart configuration.

Extracted from chart components to avoid duplication and
provide consistent chart styling across all chart types.
"""

from typing import List, Optional, Dict, Any
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.util import Pt
from pptx.dml.color import RGBColor


def configure_legend(
    chart,
    position: str = "right",
    show: bool = True,
    font_family: str = "Inter",
    font_size: int = 10
):
    """
    Configure chart legend with consistent styling.

    Args:
        chart: PowerPoint chart object
        position: Legend position (right, left, top, bottom, corner)
        show: Whether to show legend
        font_family: Font family for legend text
        font_size: Font size in points
    """
    chart.has_legend = show

    if show:
        position_map = {
            "right": XL_LEGEND_POSITION.RIGHT,
            "left": XL_LEGEND_POSITION.LEFT,
            "top": XL_LEGEND_POSITION.TOP,
            "bottom": XL_LEGEND_POSITION.BOTTOM,
            "corner": XL_LEGEND_POSITION.CORNER,
        }

        chart.legend.position = position_map.get(position.lower(), XL_LEGEND_POSITION.RIGHT)
        chart.legend.include_in_layout = False

        # Apply font styling
        if hasattr(chart.legend, 'font'):
            chart.legend.font.name = font_family
            chart.legend.font.size = Pt(font_size)


def configure_axes(
    chart,
    show_gridlines: bool = True,
    gridline_color: Optional[RGBColor] = None,
    label_font_family: str = "Inter",
    label_font_size: int = 9,
    label_color: Optional[RGBColor] = None
):
    """
    Configure chart axes with consistent styling.

    Args:
        chart: PowerPoint chart object
        show_gridlines: Whether to show gridlines
        gridline_color: Color for gridlines
        label_font_family: Font family for axis labels
        label_font_size: Font size for axis labels
        label_color: Color for axis labels
    """
    try:
        # Configure value axis
        if hasattr(chart, 'value_axis'):
            value_axis = chart.value_axis
            value_axis.has_major_gridlines = show_gridlines

            if show_gridlines and gridline_color:
                value_axis.major_gridlines.format.line.color.rgb = gridline_color
                value_axis.major_gridlines.format.line.width = Pt(0.5)

            # Format axis labels
            if hasattr(value_axis, 'tick_labels'):
                value_axis.tick_labels.font.name = label_font_family
                value_axis.tick_labels.font.size = Pt(label_font_size)
                if label_color:
                    value_axis.tick_labels.font.color.rgb = label_color

        # Configure category axis
        if hasattr(chart, 'category_axis'):
            cat_axis = chart.category_axis
            if hasattr(cat_axis, 'tick_labels'):
                cat_axis.tick_labels.font.name = label_font_family
                cat_axis.tick_labels.font.size = Pt(label_font_size)
                if label_color:
                    cat_axis.tick_labels.font.color.rgb = label_color
    except ValueError:
        # Chart type doesn't have axes (e.g., pie charts)
        pass


def set_chart_title(
    chart,
    title: str,
    font_family: str = "Inter",
    font_size: int = 16,
    font_color: Optional[RGBColor] = None,
    bold: bool = True
):
    """
    Set and style chart title consistently.

    Args:
        chart: PowerPoint chart object
        title: Chart title text
        font_family: Font family for title
        font_size: Font size in points
        font_color: Title color
        bold: Whether to make title bold
    """
    if title:
        chart.has_title = True
        chart_title = chart.chart_title
        chart_title.text_frame.text = title

        # Style the title
        para = chart_title.text_frame.paragraphs[0]
        para.font.name = font_family
        para.font.size = Pt(font_size)
        para.font.bold = bold
        if font_color:
            para.font.color.rgb = font_color


def apply_chart_colors(
    chart,
    colors: List[RGBColor]
):
    """
    Apply color palette to chart series.

    Args:
        chart: PowerPoint chart object
        colors: List of RGBColor objects for series
    """
    for idx, series in enumerate(chart.series):
        if idx < len(colors):
            fill = series.format.fill
            fill.solid()
            fill.fore_color.rgb = colors[idx]
